"""
DOCX conversion, Word document styling, email dispatching, and file cleanup.
"""

import asyncio
import os
import pathlib
import re
import smtplib
import subprocess
import logging
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import pypandoc
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from mcp_qpgenerator.generator import OUTPUT_DIR
from mcp_qpgenerator.models import SendEmailArgs

logger = logging.getLogger("mcp_qpgenerator.emailer")

FONT_NAME = "Times New Roman"

# Page layout
PAGE_WIDTH = Mm(210)        # A4
PAGE_HEIGHT = Mm(297)       # A4
MARGIN_LEFT_RIGHT = Inches(0.9)
MARGIN_TOP_BOTTOM = Inches(0.8)
ROLL_NO_BOXES = 10


def _set_font(style, size: int, bold: bool | None = None) -> None:
    """Force a font on a style and strip theme fonts/colors that override it."""
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold

    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT_NAME)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def _build_reference_docx() -> pathlib.Path:
    """Create a pandoc reference.docx with exam-paper styling and margins."""
    ref_path = OUTPUT_DIR / "qp_reference.docx"

    with open(ref_path, "wb") as f:
        subprocess.run(
            [pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"],
            stdout=f,
            check=True,
        )

    doc = Document(str(ref_path))
    styles = doc.styles

    # Page size and margins
    for section in doc.sections:
        section.page_width = PAGE_WIDTH
        section.page_height = PAGE_HEIGHT
        section.left_margin = MARGIN_LEFT_RIGHT
        section.right_margin = MARGIN_LEFT_RIGHT
        section.top_margin = MARGIN_TOP_BOTTOM
        section.bottom_margin = MARGIN_TOP_BOTTOM

    # Body text styles pandoc uses
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in styles]:
            _set_font(styles[name], 11)
            styles[name].paragraph_format.space_after = Pt(6)

    # Section headings (### in markdown): black, bold, centered
    if "Heading 3" in [s.name for s in styles]:
        h3 = styles["Heading 3"]
        _set_font(h3, 13, bold=True)
        h3.font.italic = False
        h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h3.paragraph_format.space_before = Pt(14)
        h3.paragraph_format.space_after = Pt(4)

    # Custom styles used by the markdown (matched by name)
    uni = styles.add_style("QPUniversity", WD_STYLE_TYPE.PARAGRAPH)
    uni.base_style = styles["Normal"]
    _set_font(uni, 16, bold=True)
    uni.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    uni.paragraph_format.space_after = Pt(4)

    ctr = styles.add_style("QPCenter", WD_STYLE_TYPE.PARAGRAPH)
    ctr.base_style = styles["Normal"]
    _set_font(ctr, 12)
    ctr.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ctr.paragraph_format.space_after = Pt(2)

    # Italic note under each section title
    note = styles.add_style("QPNote", WD_STYLE_TYPE.PARAGRAPH)
    note.base_style = styles["Normal"]
    _set_font(note, 11)
    note.font.italic = True
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(8)

    # Placeholder paragraph for Roll No. boxes
    roll = styles.add_style("QPRollNo", WD_STYLE_TYPE.PARAGRAPH)
    roll.base_style = styles["Normal"]
    _set_font(roll, 11)

    doc.save(str(ref_path))
    return ref_path


def _set_cell_borders(cell, sz: int = 8) -> None:
    """Draw a thin black border on all four sides of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tc_pr.append(borders)


def _build_roll_no_table(doc):
    """Return a right-aligned 'Roll No.:' label followed by empty boxes."""
    table = doc.add_table(rows=1, cols=ROLL_NO_BOXES + 1)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    table.autofit = False

    row = table.rows[0]
    row.height = Inches(0.3)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    for i, cell in enumerate(row.cells):
        cell.width = Inches(0.95) if i == 0 else Inches(0.28)
        para = cell.paragraphs[0]
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.space_before = Pt(0)
        if i == 0:
            run = para.add_run("Roll No.:")
            run.bold = True
            run.font.name = FONT_NAME
            run.font.size = Pt(11)
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_cell_borders(cell)
    return table


def _replace_roll_no_placeholder(doc) -> None:
    """Swap the QPRollNo placeholder paragraph for the real boxes."""
    for para in doc.paragraphs:
        if para.style is not None and para.style.name == "QPRollNo":
            table = _build_roll_no_table(doc)
            anchor = para._p
            anchor.addprevious(table._tbl)
            anchor.getparent().remove(anchor)
            break


def _right_align_marks(doc) -> None:
    """Move the trailing '[5]' style marks of each question to the right edge."""
    section = doc.sections[0]
    tab_pos = section.page_width - section.left_margin - section.right_margin

    for para in doc.paragraphs:
        if len(para.runs) < 2:
            continue
        last = para.runs[-1]
        if not re.fullmatch(r"\[\d+\]", last.text.strip()):
            continue

        prev = para.runs[-2]
        prev.text = prev.text.rstrip()
        para.paragraph_format.tab_stops.add_tab_stop(tab_pos, WD_TAB_ALIGNMENT.RIGHT)
        last.text = "\t" + last.text.strip()


def _postprocess_docx(docx_path: pathlib.Path) -> None:
    doc = Document(str(docx_path))
    _replace_roll_no_placeholder(doc)
    _right_align_marks(doc)
    doc.save(str(docx_path))


def convert_and_send_email_sync(
    markdown_file_path: str,
    recipients: list[str],
    subject: str,
    body: str,
) -> str:
    """Synchronous core logic for DOCX conversion, email sending, and file cleanup."""
    gmail_address = os.environ.get("GMAIL_ADDRESS")
    gmail_password = os.environ.get("GMAIL_APP_PASSWORD")
    if not gmail_address or not gmail_password:
        raise RuntimeError(
            "Missing email credentials: ensure GMAIL_ADDRESS and "
            "GMAIL_APP_PASSWORD are set in environment"
        )

    md_path = pathlib.Path(markdown_file_path)
    if not md_path.exists():
        raise FileNotFoundError(f"Markdown file not found: {markdown_file_path}")

    docx_path = md_path.with_suffix(".docx")
    try:
        ref_docx = _build_reference_docx()
        pypandoc.convert_file(
            str(md_path),
            "docx",
            format="markdown",
            outputfile=str(docx_path),
            extra_args=[f"--reference-doc={ref_docx}"],
        )
        _postprocess_docx(docx_path)

        msg = MIMEMultipart()
        msg["From"] = gmail_address
        msg["To"] = ", ".join(recipients)
        msg["Subject"] = subject
        msg.attach(MIMEText(body, "plain"))

        with open(docx_path, "rb") as f:
            part = MIMEBase("application", "octet-stream")
            part.set_payload(f.read())
        encoders.encode_base64(part)
        part.add_header(
            "Content-Disposition",
            f'attachment; filename="{docx_path.name}"',
        )
        msg.attach(part)

        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(gmail_address, gmail_password)
            server.sendmail(gmail_address, recipients, msg.as_string())
    finally:
        # Immediately delete temporary files to clean up disk space
        if md_path.exists():
            try:
                md_path.unlink()
            except Exception:
                pass
        if docx_path.exists():
            try:
                docx_path.unlink()
            except Exception:
                pass

    return f"Email sent successfully to: {', '.join(recipients)}"


from pydantic.fields import FieldInfo


async def send_email_handler(
    markdown_file_path: str,
    recipients: list[str],
    subject: str = "Question Paper",
    body: str = "Please find the attached question paper.",
    user_confirmed: bool = False,
) -> str:
    """Async wrapper that validates Pydantic guardrails and executes email dispatch in a thread."""
    # Resolve values if invoked directly in Python where defaults might resolve to FieldInfo
    resolved_subject = "Question Paper" if isinstance(subject, FieldInfo) else subject
    resolved_body = "Please find the attached question paper." if isinstance(body, FieldInfo) else body
    resolved_confirmed = False if isinstance(user_confirmed, FieldInfo) else user_confirmed

    # Enforce Pydantic schema validation
    args = SendEmailArgs(
        markdown_file_path=markdown_file_path,
        recipients=recipients,
        subject=resolved_subject,
        body=resolved_body,
        user_confirmed=resolved_confirmed,
    )

    if args.user_confirmed is not True:
        raise ValueError(
            "Email dispatch rejected: User confirmation has not been explicitly granted. "
            "You must present the preview to the user and obtain their confirmation before setting user_confirmed=True."
        )

    clean_recipients = [r.strip() for r in args.recipients if r and r.strip()]
    if not clean_recipients:
        raise ValueError("No recipients provided. Please confirm the recipient email address with the user.")

    try:
        result = await asyncio.to_thread(
            convert_and_send_email_sync,
            args.markdown_file_path,
            clean_recipients,
            args.subject,
            args.body,
        )
        return result
    except Exception as e:
        logger.exception("Failed to convert and send email asynchronously")
        raise RuntimeError(f"Email sending failed: {e}") from e
